
import os
import logging
import traceback
import asyncio
import inspect
import re
from dataclasses import dataclass
from typing import Dict, List, Tuple, Optional
from difflib import get_close_matches
from docx import Document

from utils.document_loader import DocumentLoader
from threading import Lock
from langchain_utils.section_checker import SectionChecker
from utils.parsers import split_recommendations

from utils.docx_writer import (
    build_replacements,
    replace_placeholders_in_doc,
    insert_recommendations,
    fill_comparison_table,
    save_with_timestamp,
)

from lv_parser import split_leaflet_sections_simple
from ohlp_parser import split_ohlp_sections

# ----------------------------
# Module logger
# ----------------------------
logger = logging.getLogger(__name__)

# ----------------------------
# Caching for recommendations
# ----------------------------
_recommendation_cache: Dict[str, str] = {}
_cache_lock = Lock()


def load_cached_recommendations(rec_path: str) -> str:
    """Thread-safe cached file loader for recommendations text."""
    with _cache_lock:
        if rec_path not in _recommendation_cache:
            _recommendation_cache[rec_path] = DocumentLoader(rec_path).simple_load()
            logger.info("Загружаем рекомендации в кэш")
    logger.info("Используем кэшированный файл рекомендаций")
    return _recommendation_cache[rec_path]


# ----------------------------
# Helpers for section name cleaning & extraction
# ----------------------------

def _clean_section_name(s: str) -> str:
    """
    Очищает строку: убирает пробелы, звёздочки, кавычки и пр. мусор,
    но сохраняет <...> и, при необходимости, восстанавливает недостающую '<'.
    """
    if not s:
        return ""
    s = s.replace("\u00A0", " ")
    s = re.sub(r'\*+', '', s)
    s = re.sub(r'\s+', ' ', s)
    s = re.sub(r'^[\s\.\-–—«»"\'№]+', '', s)
    s = re.sub(r'[\s\.\-–—«»"\'№]+$', '', s)
    if s.endswith('>') and not s.startswith('<'):
        s = '<' + s
    if s.startswith('<') and '>' not in s[1:]:
        s = s + '>'
    return s.strip()


def extract_sections_from_recommendations(text: str) -> List[str]:
    """
    Разбивает текст рекомендаций по маркеру %split% и возвращает список
    заголовков (первая непустая строка каждого блока), очищенных от
    спецсимволов. Пропускает служебные секции и маркер %extra%.
    """
    parts = re.split(r'%\s*split\s*%', text or '', flags=re.IGNORECASE)
    sections: List[str] = []
    for part in parts:
        block = (part or '').strip()
        if not block:
            continue
        if block.lower().startswith('рекомендации по составлению проекта общей характеристики'):
            continue
        if block.strip().lower() == '%extra%':
            continue
        lines = block.splitlines()
        first = ''
        for ln in lines:
            if ln.strip():
                first = ln.strip()
                break
        if not first:
            continue
        clean = _clean_section_name(first)
        if clean:
            sections.append(clean)
    return sections


# ----------------------------
# Data containers
# ----------------------------

@dataclass
class LoadedDocs:
    loader_test: DocumentLoader
    loader_ref: DocumentLoader
    test_text: str
    ref_text: str
    rec_text: str


# ----------------------------
# Core steps used by report generation
# ----------------------------

async def load_texts(
    test_path: str,
    ref_path: str,
    rec_path: str,
) -> LoadedDocs:
    """Load texts concurrently in threads; recommendations via cache."""
    logger.info("Загрузка текстов…")
    loader_test = DocumentLoader(test_path)
    loader_ref = DocumentLoader(ref_path)

    test_text, ref_text = await asyncio.gather(
        asyncio.to_thread(loader_test.load),
        asyncio.to_thread(loader_ref.load),
    )
    rec_text = await asyncio.to_thread(load_cached_recommendations, rec_path)

    logger.info(
        f"TEST: {len(test_text)} chars, type={loader_test.doc_type}, drug={loader_test.drug_name!r}"
    )
    logger.info(
        f"REF:  {len(ref_text)} chars, type={loader_ref.doc_type}, drug={loader_ref.drug_name!r}"
    )
    logger.info(f"REC:  {len(rec_text)} chars")

    return LoadedDocs(loader_test, loader_ref, test_text, ref_text, rec_text)


async def extract_sections(rec_text: str) -> List[str]:
    logger.info("Извлечение разделов из рекомендаций…")
    return await asyncio.to_thread(extract_sections_from_recommendations, rec_text)


async def split_into_sections(
    *,
    doc_type: str,
    test_text: str,
    ref_text: str,
    rec_text: str,
    sections: List[str],
) -> Tuple[Dict[str, str], Dict[str, str], Dict[str, str]]:
    """Split test/ref/rec texts into section blocks concurrently."""
    logger.info("Разбиение текстов по разделам…")
    if doc_type == "leaflet":
        test_blocks, ref_blocks = await asyncio.gather(
            asyncio.to_thread(split_leaflet_sections_simple, test_text, sections),
            asyncio.to_thread(split_leaflet_sections_simple, ref_text, sections),
        )
    else:
        test_blocks, ref_blocks = await asyncio.gather(
            asyncio.to_thread(split_ohlp_sections, test_text, sections),
            asyncio.to_thread(split_ohlp_sections, ref_text, sections),
        )
    recs_blocks = await asyncio.to_thread(split_recommendations, rec_text, sections)
    return test_blocks, ref_blocks, recs_blocks


async def check_recommendations(
    *,
    checker: SectionChecker,
    sections: List[str],
    test_blocks: Dict[str, str],
    recs_blocks: Dict[str, str],
    concurrency: int = 5,
) -> Dict[str, dict]:
    """Check recommendations per section with bounded concurrency."""
    logger.info("Проверка рекомендаций по разделам…")
    sem = asyncio.Semaphore(concurrency)

    async def check_one(sec: str):
        actual_text = test_blocks.get(sec, "")

        # ищем самое близкое название среди ключей recs_blocks
        matches = get_close_matches(sec, recs_blocks.keys(), n=1, cutoff=0.7)
        if not matches:
            return sec, {'compliance': 'not_complied', 'comments': ''}

        rec_part = recs_blocks[matches[0]]
        if not rec_part.strip():
            return sec, {'compliance': 'not_complied', 'comments': ''}

        async with sem:
            try:
                # ⬇️ ВСЕГДА используем асинхронный LLM-вызов
                result = await checker.check_recommends_async(rec_part, actual_text)
                return sec, result
            except Exception as e:
                logger.error(f"Ошибка при проверке рекомендаций для '{sec}': {e}")
                return sec, {'compliance': 'not_complied', 'comments': f'Ошибка: {e}'}

    pairs = await asyncio.gather(*(check_one(sec) for sec in sections))
    return {sec: res for sec, res in pairs}


def load_template(template_dir: str, template_name: str) -> Document:
    logger.info("Загрузка шаблона DOCX…")
    template_path = os.path.join(template_dir, template_name)
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон не найден: {template_path}")
    return Document(template_path)


def write_recommendations_and_table(
    doc: Document,
    *,
    recommendations: Dict[str, dict],
    sections: List[str],
    ref_blocks: Dict[str, str],
    test_blocks: Dict[str, str],
    table_index: int = 2,
) -> None:
    logger.info("Вставка рекомендаций и заполнение таблицы…")
    
    try:
        insert_recommendations(doc, recommendations)
    except Exception as e:
        logger.error(f"Ошибка вставки рекомендаций: {e}")

    try:
        fill_comparison_table(
            doc,
            recommended_sections=sections,
            ref_blocks=ref_blocks,
            test_blocks=test_blocks,
            table_index=table_index
        )
    except Exception as e:
        logger.error(f"Ошибка заполнения таблицы сравнения: {e}")


def replace_placeholders(doc: Document, *, drug_name_ref: str, drug_name_test: str) -> None:
    logger.info("Замена плейсхолдеров…")
    try:
        replacements = build_replacements(drug_name_ref, drug_name_test)
        replace_placeholders_in_doc(doc, replacements)
    except Exception as e:
        logger.error(f"Ошибка замены плейсхолдеров: {e}")


def save_report(
    doc: Document,
    *,
    filetype: str,
    original_filename: str,
    output_dir: str,
    prefix: str,
) -> str:
    logger.info("Сохранение итогового документа…")
    output_path = save_with_timestamp(
        doc,
        filetype=filetype,
        original_filename=original_filename,
        output_dir=output_dir,
        prefix=prefix,
    )
    logger.info(f"Отчёт сохранён: {output_path}")
    return output_path


# ----------------------------
# Main async API
# ----------------------------

async def generate_report_async(
    test_path: str,
    ref_path: str,
    rec_path: str,
    template_name: str,
    template_dir: str,
    output_dir: str,
    provider: str,
    prefix: str,
    *,
    concurrency: int = 5,
    test_blocks: Optional[Dict[str, str]] = None,
    ref_blocks: Optional[Dict[str, str]] = None,
    sections_order: Optional[List[str]] = None,  # можно пока не использовать
) -> str:
    logger.info("=== Начало генерации отчёта ===")
    try:
        docs = await load_texts(test_path, ref_path, rec_path)

        # эталонные секции берём из рекомендаций (как и было)
        sections = await extract_sections(docs.rec_text)
        doc = load_template(template_dir, template_name)

        # рекомендации режем всегда
        recs_blocks = await asyncio.to_thread(split_recommendations, docs.rec_text, sections)

        # test/ref блоки: либо пришли из UI, либо режем автоматически (как и было)
        if test_blocks is None or ref_blocks is None:
            auto_test_blocks, auto_ref_blocks, _ = await split_into_sections(
                doc_type=docs.loader_test.doc_type,
                test_text=docs.test_text,
                ref_text=docs.ref_text,
                rec_text=docs.rec_text,
                sections=sections,
            )
            test_blocks = test_blocks or auto_test_blocks
            ref_blocks = ref_blocks or auto_ref_blocks

        checker = SectionChecker(api_provider=provider)
        recommendations = await check_recommendations(
            checker=checker,
            sections=sections,
            test_blocks=test_blocks,
            recs_blocks=recs_blocks,
            concurrency=concurrency,
        )

        write_recommendations_and_table(
            doc,
            recommendations=recommendations,
            sections=sections,
            ref_blocks=ref_blocks,
            test_blocks=test_blocks,
            table_index=2,
        )

        replace_placeholders(
            doc,
            drug_name_ref=docs.loader_ref.drug_name,
            drug_name_test=docs.loader_test.drug_name,
        )

        original_filename = os.path.splitext(os.path.basename(docs.loader_test.file_path))[0]
        output_path = save_report(
            doc,
            filetype=docs.loader_test.doc_type,
            original_filename=original_filename,
            output_dir=output_dir,
            prefix=prefix,
        )
        logger.info("=== Генерация отчёта завершена ===")
        return output_path

    except Exception as e:
        logger.critical(f"Невозможно сгенерировать отчёт: {e}")
        logger.debug(traceback.format_exc())
        raise

